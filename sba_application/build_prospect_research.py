"""Prospect research deliverable: Tyler TX investor list for the $500K note raise.

Combines two research tracks (angel/HNW + physicians) into a ranked docx and CSV.
Research date 2026-07-16; public-web sources only. Output:
sba_application/12_investor_raise/Prospect_Research_Tyler.{docx,csv}
"""
import csv
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

OUT = "sba_application/12_investor_raise/"

# ---- Track 1: angel / HNW / business (fields: name, role, why, path, tier, src, ver)
ANGELS = [
    ("Raymond C. McKinney, CPA",
     "CEO & President, The Genecov Group (Tyler's largest commercial landlord + oil/gas); director, Southside Bancshares; UT Health Tyler Board of Trustees; past chairman, Tyler Economic Development Council",
     "Sits at the intersection of every target category: real estate/oil wealth, bank board, hospital trustee, economic development. CPA - can underwrite a note quickly.",
     "Banking (strongest Jim Bullard fit) - Southside Bank board circle; alternates: TEDC, UT Health Tyler board",
     "$100K-$250K",
     "https://investors.southside.com/news/news-details/2025/Southside-Bancshares-Inc--Announces-New-Directors-and-Chief-Operating-Officer/default.aspx",
     "Search-verified"),
    ("John R. 'Bob' Garrett",
     "Vice Chairman, Southside Bancshares; President, Fair Oil Company since 2002; residential & commercial real estate developer",
     "Triple-category: oil/gas family enterprise (Fair Oil / R.W. Fair legacy), commercial RE, senior bank board seat. Long Tyler civic tenure.",
     "Banking - Southside board via Bullard's bank relationships",
     "$250K+",
     "https://www.sec.gov/Archives/edgar/data/705432/000070543218000031/proxy12-31x17.htm",
     "Search-verified (SEC proxy excerpt)"),
    ("Jeff Austin III",
     "Chairman of the Board, Austin Bank (branches across Smith County); 4th-generation East Texas banker",
     "Family owns Austin Bank outright; personal wealth plus a bank whose lending footprint covers Tyler. Active in statewide banking circles.",
     "Banking - IBAT/Texas banking circles; Austin Bank commercial relationships",
     "$100K-$250K",
     "https://www.austinbank.com/our-story/our-history",
     "Search-verified"),
    ("James I. Perkins",
     "Chairman & President, Citizens 1st Bank (Tyler); 3rd-generation banking family",
     "Owner-operator of a private Tyler bank; family recognized with 2023 Texas Higher Education Distinguished Service Award - demonstrated large-gift philanthropy.",
     "Banking - community bank chairman circles (IBAT/TBA), direct fit for Bullard",
     "$100K-$250K",
     "https://member.texasbankers.com/Magazine/magazine/Features/2023-07/23-07-Spotlight-James-Perkins.aspx",
     "Search-verified"),
    ("Rogers Pope Jr.",
     "Chairman & CEO, Texas Bank and Trust (Longview; major Tyler presence); past chairman, Independent Bankers Association of Texas (2015-16)",
     "Controls one of East Texas's largest family-run banks; IBAT chairmanship = deep statewide banker network.",
     "Banking - IBAT network; TB&T Tyler advisory circle",
     "$100K-$250K",
     "https://bankersdigest.com/texas-bank-and-trust-announces-chairmanship-transition/",
     "Search-verified"),
    ("Russ Gideon",
     "COO, Austin Bank; Director, CHRISTUS Trinity Mother Frances Foundation board",
     "Rare dual node: senior bank executive AND hospital-foundation director - banking warm path plus demonstrated Tyler healthcare philanthropy.",
     "Banking (Austin Bank via Bullard) or CHRISTUS foundation board",
     "$50K-$100K",
     "https://www.christushealth.org/connect/donate/trinity-mother-frances/team",
     "Search-verified"),
    ("Dr. Kirk A. Calhoun",
     "Director, Southside Bancshares; President Emeritus, UT Tyler / UT Health Science Center; physician (nephrology, non-practicing)",
     "Built UT Tyler's medical enterprise over two decades; now on a Tyler bank board. Best clinical-credibility validator an East Texas hospice raise could have. Now president of UNT Health Fort Worth - confirm continued Tyler presence.",
     "Banking - Southside board; alternates: UT Tyler leadership alumni, W. Fair Foundation circle",
     "$50K-$100K",
     "https://bankersdigest.com/calhoun-former-university-of-texas-tyler-president-elected-to-southside-bancshares-board/",
     "Search-verified"),
    ("Joe Cavender (with brothers Mike & Clay)",
     "President, Cavender's Boot City (Tyler HQ, 100+ stores in 15 states; family-owned)",
     "One of Tyler's wealthiest operating families; company sits on the TEDC board; long civic sponsorship record. Private family - can move quickly on a $250K note.",
     "Chamber/TEDC circles; Tyler retail/business community",
     "$250K+",
     "https://www.tylertodaymagazine.com/local/cavenders-boot-city-celebrating-60-years/",
     "Search-verified"),
    ("Soules family - Mark Soules & John Soules Jr. (Co-CEOs); John Soules Sr. (founder)",
     "Soules Foods, Tyler HQ; nation's leading fajita processor, 4th-generation family leadership",
     "Major private food-manufacturing wealth; a John Soules sits on the CHRISTUS Trinity Mother Frances Foundation board - direct evidence of healthcare philanthropy.",
     "CHRISTUS TMF Foundation board; Tyler Chamber; UT Tyler events",
     "$100K-$250K",
     "https://www.soulesfoods.com/our-story/",
     "Search-verified"),
    ("Brad Brookshire (with Ann Warmack Brookshire)",
     "Retired Chairman & CEO, Brookshire Grocery Co. (Tyler HQ, 200+ stores); retired March 2026 after 49 years",
     "Family net worth ~$1B (Forbes-listed); just exited CEO role - classic legacy-phase profile; SMU/TJC scholarship philanthropy. More philanthropist than note investor, but unmatched local capacity.",
     "Tyler Chamber (Brookshire's is an anchor member); TJC/SMU philanthropy circles; church networks (specific congregation unverified)",
     "$250K+ (if interested at all)",
     "https://www.ketk.com/news/local-news/brookshire-grocery-ceo-retiring-after-49-years-in-familys-business/",
     "Search-verified"),
    ("Robert Peltier family",
     "Founder/owner, Peltier Enterprises - 6 auto dealerships (Chevrolet, Nissan, Kia, Subaru) in Tyler/Longview",
     "35+ year dealer group; gives $50K/yr in scholarships - $50K is literally their annual giving unit. Dealers understand cash-flowing local service businesses.",
     "Tyler Chamber; East Texas auto/commercial banking relationships (floorplan lenders may connect to Bullard)",
     "$50K-$100K",
     "https://www.peltier.net/dealership/about",
     "Search-verified"),
    ("Sam Roosth & Steve Roosth",
     "Sam: Roosth Production Co. (oil & gas, Tyler); Steve: family real estate arm; heirs to Roosth & Genecov oil fortune",
     "Century-old Tyler oil/land family; endowed a UT Tyler chair in 2000; royalty-style periodic income fits the quarterly-interest note mental model.",
     "Genecov Group connection (Raymond McKinney manages the affiliated side); Tyler oil circles; Jewish community philanthropy network",
     "$50K-$100K each",
     "https://business.tylertexas.com/list/member/roosth-production-co-3518",
     "Search-verified"),
    ("Robyn Rogers",
     "President, The Robert M. Rogers Foundation (Tyler; $34-52M assets; $10M gift to UT Tyler School of Medicine, 2022)",
     "Foundation focus (East Texas behavioral health) is mission-adjacent to hospice. Foundation money would be grants, not notes - target personal capacity and her endorsement network.",
     "UT Tyler School of Medicine / UT Health leadership; Philanthropy Southwest",
     "Philanthropic ally; personal check $50K (appetite unverified)",
     "https://www.uthct.edu/the-robert-m-rogers-foundation-gifts-10-million-to-the-ut-tyler-school-of-medicine/news/2022/07/18/34647/",
     "Search-verified"),
    ("Mewbourne family (current principals unverified)",
     "Owners (100%), Mewbourne Oil Co. - Tyler-HQ'd, largest private oil producer in the U.S.; founder Curtis Mewbourne died June 2026",
     "Deepest capacity in Tyler by far; estate/legacy transition may open family-office-style allocations. Access is hard; only pursue with a genuine mutual contact.",
     "Tyler oil & gas circles; UT-system engineering philanthropy",
     "$250K+ (access-limited)",
     "https://www.mewbourne.com/meet-our-founder/",
     "Search-verified"),
    ("Randall Childress",
     "Chair, CHRISTUS Trinity Mother Frances Foundation board; Edward Jones financial advisor",
     "Chairs the region's flagship hospital foundation - the best door-opener into Tyler's healthcare-philanthropy donor base; as a wealth advisor can also involve clients (506(b): pre-existing relationships only).",
     "CHRISTUS TMF Foundation; Edward Jones/wealth-management circles",
     "$50K (connector value > check)",
     "https://www.christushealth.org/connect/donate/trinity-mother-frances/team",
     "Search-verified"),
    ("Tom Perkins",
     "2025-26 Board Chair, Tyler Area Chamber of Commerce; Edward Jones financial advisor",
     "Current chamber chair = fastest legitimizer for a new Tyler healthcare employer; advisor practice means HNW client network.",
     "Tyler Chamber (direct)",
     "$50K (connector)",
     "https://tylerpaper.com/2025/10/09/tyler-area-chamber-of-commerce-annual-meeting-celebrates-125-years-awarding-dedicated-citizens/",
     "Search-verified"),
    ("Terry Giles",
     "Franchise owner, Dairy Queens of East Texas; Director, CHRISTUS TMF Foundation",
     "Multi-unit franchise wealth plus proven healthcare-philanthropy engagement.",
     "CHRISTUS TMF Foundation board",
     "$50K-$100K",
     "https://www.christushealth.org/connect/donate/trinity-mother-frances/team",
     "Search-verified"),
    ("Steve Hellmuth",
     "Owner, Dakota's Restaurant (Tyler); Director, CHRISTUS TMF Foundation",
     "Local operating-business owner with hospital-foundation seat; plausible smallest-tier note.",
     "CHRISTUS TMF Foundation board",
     "$50K",
     "https://www.christushealth.org/connect/donate/trinity-mother-frances/team",
     "Search-verified"),
    ("Hibbs-Hallmark & Company principals",
     "Family-owned Tyler insurance agency (one of the largest independent agencies in East Texas)",
     "Named in a local wealthiest-Tyler-families roundup; insurance principals understand actuarial/healthcare risk. Confirm current ownership before outreach.",
     "Tyler Chamber; commercial insurance / banking referral circles",
     "$50K-$100K (guess)",
     "https://knue.com/richest-families-tyler-tx/",
     "UNVERIFIED (weak source)"),
    ("Jeb W. Jones - CONFLICT FLAG",
     "CEO, Pro Star Rental / President & GP, Pro Star Capital LP; new Southside Bancshares director; Vice Chair, CHRISTUS Health Northeast Texas Region; PAST CHAIR, HOSPICE OF EAST TEXAS",
     "On paper the perfect prospect - but he is past board chair of the incumbent nonprofit hospice you'd compete with. Assume loyalty conflict and information risk. Approach only deliberately, if at all.",
     "Banking (Southside) - but see conflict flag",
     "N/A until conflict assessed",
     "https://investors.southside.com/news/news-details/2025/Southside-Bancshares-Inc--Announces-New-Directors-and-Chief-Operating-Officer/default.aspx",
     "Search-verified"),
]

# ---- Track 2: physicians (fields: name, role, why, aks, path, tier, src, ver)
PHYSICIANS = [
    ("Dr. Charles R. Gordon",
     "Neurosurgery (spine); founder, Texas Spine & Joint Hospital; founder, Precision Spine Care; founder, Flexus Spine (device co.)",
     "Known lead - confirmed serial physician-entrepreneur: hospital, multi-site practice, and device company.",
     "LOW - spine surgery; does not refer to hospice",
     "Direct (existing relationship)",
     "$100K-$250K",
     "https://tsjh.org/charles-r-gordon-m-d/",
     "Search-verified"),
    ("Dr. Michael E. Russell II",
     "Orthopedic spine surgery, Azalea Orthopedics; co-founder & past board chairman, TSJH; past president, Physician Hospitals of America",
     "National physician-ownership advocate; hospital equity; practice leadership - arguably the strongest capacity+ideology fit in Tyler for a physician-owned venture.",
     "LOW - orthopedics",
     "Via Dr. Gordon (TSJH co-founders) or Azalea network",
     "$100K-$250K",
     "https://azaleaortho.com/doctor/michael-e-russell-ii-md/",
     "Search-verified"),
    ("Dr. Kevin Pauza",
     "Physiatry / interventional spine; co-founder & founding partner, TSJH; inventor of the Discseel procedure",
     "Hospital co-founder, procedure/IP entrepreneur. DILIGENCE FLAG (unverified): a tylerpaper.com headline about a Tyler doctor disciplined by the medical board surfaced in his search results - confirm whether it concerns him before approach.",
     "LOW - interventional spine/PM&R",
     "Via Dr. Gordon / TSJH founders' circle",
     "$100K-$250K",
     "https://en.wikipedia.org/wiki/Kevin_Pauza",
     "Search-verified; diligence flag"),
    ("Dr. Aaron K. Calodney",
     "Anesthesiology / interventional pain, Precision Spine Care + TSJH; past president, Texas Pain Society; ASIPP board",
     "National pain-medicine leader, TSJH privileges since 1991, likely TSJH ownership stake; industry consulting income.",
     "LOW - anesthesia/pain",
     "Via Precision Spine Care partners (Gordon) or Texas Pain Society",
     "$50K-$100K",
     "https://tsjh.org/aaron-k-calodney-m-d/",
     "Search-verified"),
    ("Dr. John Priddy",
     "Orthopedic surgery (foot/ankle), Azalea Orthopedics; clinic president 2022-2025; TSJH-affiliated",
     "Just completed practice-president term - governance experience, partner-level equity, freed-up bandwidth.",
     "LOW - orthopedics",
     "Azalea partners via Dr. Russell; TSJH network",
     "$50K-$100K",
     "https://azaleaortho.com/doctor/john-priddy-md/",
     "Search-verified"),
    ("Dr. Patrick Wupperman",
     "Orthopedic surgery / sports medicine (Andrews fellowship), Azalea Orthopedics; incoming practice president, 2026",
     "Rising leader of the region's dominant ortho group; partner equity.",
     "LOW - orthopedics/sports medicine",
     "Azalea network via Russell/Priddy",
     "$50K",
     "https://azaleaortho.com/dr-wupperman-named-azalea-president/",
     "Search-verified"),
    ("Dr. Leo W. Mack Jr.",
     "Ophthalmology (cataract); private practice in Tyler since 1976; operates at NovaMed Surgery Center of Tyler",
     "~50-year career, likely retired or winding down; ASC operating history suggests ownership-mindedness and accumulated capital.",
     "LOW - ophthalmology; retired/near-retired",
     "NovaMed ASC physician group; SCMS senior-member circle",
     "$50K",
     "https://novamedsurgerycenteroftyler.com/physicians/leo-w-mack-jr-md",
     "Search-verified"),
    ("Dr. Michael C. Ford",
     "Ophthalmology, Tyler; operates at NovaMed Surgery Center; CHRISTUS + UT Health affiliations",
     "25+ year cataract/glaucoma practice; ASC user (ownership unverified).",
     "LOW - ophthalmology",
     "NovaMed ASC group; Dr. Mack",
     "$50K",
     "https://novamedsurgerycenteroftyler.com/physicians/michael-ford-md",
     "Search-verified"),
    ("Dr. Craig E. Harrison",
     "Plastic surgery, longtime private practice, Tyler",
     "Established cosmetic practice = cash-pay income; ownership of surgical suite unverified.",
     "LOW - plastic surgery",
     "Chamber of Commerce; SCMS",
     "$50K",
     "https://taabsink.wixsite.com/plasticsurgerytyler/dr-harrison",
     "Search-verified"),
    ("Dr. Mark Wallis (and Dr. Luke Wallis)",
     "Dermatology; Wallis Dermatology - multi-site private group (Tyler, Longview, Marshall)",
     "Family-owned multi-location practice that has stayed independent of PE roll-ups; business-builder profile.",
     "LOW - dermatology",
     "SCMS; East Texas referral community; chamber",
     "$50K",
     "https://www.wallisderm.com/",
     "Search-verified"),
    ("Dr. Noah Israel",
     "Cardiology; founded Cardiovascular Associates of East Texas (1982), grew into CardiaStream, sold to CHRISTUS Health in 2023",
     "Built and exited the largest cardiology group in East Texas - proven liquidity event; likely retired or senior. Confirm current status post-acquisition.",
     "HIGH - cardiology plausibly refers to hospice; requires regulatory counsel clearance (mitigated if fully retired - verify)",
     "CardiaStream alumni network; CHRISTUS medical staff; SCMS",
     "$100K+",
     "https://www.christushealth.org/connect/news/christus-health-acquires-cardiastream",
     "Search-verified"),
    ("Dr. C. Fagg Sanford",
     "Cardiology; co-built the first heart program at Mother Frances Hospital; CardiaStream principal",
     "Same exit as Dr. Israel; senior/likely retired. Status unverified.",
     "HIGH - cardiology; counsel clearance required",
     "Via Dr. Israel / CHRISTUS heart program alumni",
     "$50K-$100K",
     "https://www.christushealth.org/connect/news/christus-health-acquires-cardiastream",
     "UNVERIFIED status"),
    ("Dr. Steven P. Keuer",
     "Internal medicine; president, CHRISTUS Trinity Clinic & CMO, CHRISTUS Trinity Mother Frances Health System",
     "Top physician executive of Tyler's largest system; convenes the Trinity Clinic physician network.",
     "HIGH - IM + controls a primary-care/hospitalist network that is a major hospice referral source; counsel clearance essential",
     "CHRISTUS medical staff leadership; SCMS",
     "$50K",
     "https://www.christushealth.org/trinity/clinic/find-a-doctor/steven-keuer",
     "Search-verified"),
    ("Dr. Suman Sinha",
     "Pulmonology/critical care; Chief of Pulmonary Medicine, CHRISTUS Health, Tyler; Texas Medical Association physician leader",
     "Department chief; TMA leadership = wide physician network; politically engaged.",
     "HIGH - pulmonology/critical care refers directly to hospice; counsel clearance required",
     "TMA/SCMS leadership track",
     "$50K",
     "https://www.linkedin.com/in/suman-sinha-a0290a1b7/",
     "Search-verified"),
    ("Dr. Mike Lamanteer",
     "Internal medicine/hospitalist; CMO, UT Health East Texas (since Sept 2023)",
     "System CMO of Tyler's other major system.",
     "HIGH - hospitalist/IM background + oversees discharge pathways; counsel clearance essential",
     "UT Health East Texas medical staff; UT Tyler School of Medicine circles",
     "$50K",
     "https://theorg.com/org/ut-health-east-texas/org-chart/mike-lamanteer-m-d",
     "Search-verified"),
    ("Dr. Kyna Schreiber",
     "Vascular neurology, UT Health East Texas; assistant professor, UT Tyler; reported president, Smith County Medical Society",
     "If SCMS president, she is the gateway to the county's organized-medicine network. SCMS role from search summary only.",
     "HIGH - neurology (stroke/dementia) refers to hospice; counsel clearance required",
     "SCMS itself",
     "$50K",
     "https://www.uttyler.edu/directory/neurology-residency/kyna-schreiber.php",
     "UNVERIFIED (SCMS role)"),
]

TOP10 = [
    ("1", "Dr. Charles R. Gordon", "Physician", "Direct relationship; anchor for the whole TSJH/spine cluster", "$100K-$250K"),
    ("2", "Raymond C. McKinney", "Angel/HNW", "Southside board + Genecov + UT Health trustee; Bullard banking path", "$100K-$250K"),
    ("3", "John R. 'Bob' Garrett", "Angel/HNW", "Southside vice chair + Fair Oil; largest single-check potential on the bank path", "$250K+"),
    ("4", "Dr. Michael E. Russell II", "Physician", "TSJH co-founder, physician-ownership advocate; one intro from Gordon", "$100K-$250K"),
    ("5", "Jeff Austin III", "Angel/HNW", "Austin Bank owner-chairman; IBAT circles = Bullard's lane", "$100K-$250K"),
    ("6", "James I. Perkins", "Angel/HNW", "Citizens 1st Bank owner; single-family decision-maker", "$100K-$250K"),
    ("7", "Rogers Pope Jr.", "Angel/HNW", "Texas Bank and Trust; IBAT past chairman network", "$100K-$250K"),
    ("8", "Joe Cavender", "Angel/HNW", "Cavender's Boot City family; fast-moving private capital", "$250K+"),
    ("9", "Soules family", "Angel/HNW", "Soules Foods + a seat on the CHRISTUS TMF Foundation board", "$100K-$250K"),
    ("10", "Dr. Kevin Pauza", "Physician", "TSJH co-founder via Gordon; clear the diligence flag first", "$100K-$250K"),
]


def _style(doc):
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)


def _h(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x3B, 0x2D)
    return h


def _p(doc, text, bold=False, size=10):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    return p


def _table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, htxt in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = htxt
        for par in cell.paragraphs:
            for run in par.runs:
                run.bold = True
                run.font.size = Pt(8)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
            for par in cells[i].paragraphs:
                for run in par.runs:
                    run.font.size = Pt(8)
    return t


def build_docx():
    doc = Document()
    _style(doc)

    title = doc.add_heading("Investor Prospect Research - Tyler, Texas", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("$500K promissory-note raise | Private offering (Reg D 506(b), warm introductions only) | Research date: July 16, 2026")
    r.font.size = Pt(9)
    r.italic = True

    _h(doc, "How to read this list")
    _p(doc, "Every entry came from public sources (news, org websites, filings, public profiles). Most Tyler-area sites block "
            "automated access, so entries were verified through live search results quoting those pages; titles and roles should be "
            "re-confirmed before outreach. Entries marked UNVERIFIED rest on a single source. Nothing here is contact-list material - "
            "this is a warm-introduction map. Under 506(b) there is no general solicitation: every approach goes through a "
            "pre-existing relationship chain.")

    _h(doc, "Top 10 overall")
    _table(doc, ["#", "Name", "Track", "Why they lead the list", "Est. tier"], TOP10)

    _h(doc, "Where the warm paths concentrate")
    for x in [
        "1. The Southside Bank boardroom - four prospects on one board (McKinney, Garrett, Dr. Calhoun, plus the flagged Jeb Jones). "
        "If Jim Bullard's banking relationships reach anyone at Southside, one introduction cascades to three-plus qualified prospects. Work this first.",
        "2. The TSJH / Precision Spine Care / Azalea triangle - Dr. Gordon is the hub. One anchor commitment from him plausibly unlocks "
        "4-6 LOW-risk physician checks from people who have already syndicated a healthcare facility together.",
        "3. The CHRISTUS Trinity Mother Frances Foundation board - a pre-assembled roster of Tyler HNW individuals already committed to "
        "healthcare philanthropy (Childress, Gideon, Giles, Hellmuth, a Soules).",
        "4. The three family bank dynasties (Austin, Perkins, Pope) - single-family decision-makers reachable through IBAT/Texas Bankers "
        "Association circles, again Bullard's lane.",
        "5. Genecov Group is a hidden hub - McKinney links the Genecov family, the Roosth oil family, Southside's board, TEDC, and UT Health Tyler's trustees.",
        "6. Chamber/TEDC paths (Tom Perkins, Cavender, Peltier, Brookshire) are broader but slower - use for credibility events, not first checks.",
    ]:
        _p(doc, x)

    _h(doc, "Track 1 - Angel / high-net-worth / business prospects (ranked)")
    _table(doc, ["#", "Name", "Role / affiliation", "Why them", "Warm path", "Est. tier", "Verification"],
           [(i + 1, n, role, why, path, tier, ver) for i, (n, role, why, path, tier, src, ver) in enumerate(ANGELS)])

    _h(doc, "Track 2 - Physician prospects (ranked; LOW referral risk ranks above HIGH)")
    _p(doc, "AKS tag: paying note interest to a physician who refers patients to the hospice is remuneration to a referral source under "
            "the federal Anti-Kickback Statute. LOW = retired/non-referring specialty. HIGH = plausibly refers; healthcare regulatory "
            "counsel must clear each HIGH prospect in writing (or confirm full retirement) before accepting funds.", bold=True)
    _table(doc, ["#", "Name", "Specialty / affiliation", "Why them", "AKS tag", "Warm path", "Est. tier", "Verification"],
           [(i + 1, n, role, why, aks, path, tier, ver) for i, (n, role, why, aks, path, tier, src, ver) in enumerate(PHYSICIANS)])

    _h(doc, "Additional physician leads (single-source, not ranked)")
    for x in [
        "Smith County Medical Society officers (reported; SCMS site not loadable): Dr. Jennifer Newton (VP), Dr. Ryan Menard (Secretary), Dr. Lisa Allen (Treasurer).",
        "Precision Spine Care partners - all LOW risk, all reachable through the Gordon/TSJH channel: Drs. William Dreiss, Duane Griffith, Jonathan Blau, Austin Harper, Ellisiv Lien, Ameer Ali.",
        "Tyler Radiology Associates, P.A. - independent radiologist-owned group (LOW risk); individual partner names not yet captured.",
    ]:
        _p(doc, "  -  " + x)

    _h(doc, "Organizations checked (not individual prospects)")
    _table(doc, ["Organization", "Finding"], [
        ("Tyler Texas Angel Network", "Only formal angel group found for Tyler; its website no longer resolves - likely inactive. Verify before relying on it."),
        ("East Texas Angel Network (Longview)", "NOT an investor group - it is Neal McCoy's charity for children with terminal illnesses. Do not pitch; potential goodwill partner at most."),
        ("R.W. Fair Foundation", "Fair-family grantmaking foundation; health & human services is a stated grant area. Grant prospect, not note investor; Bob Garrett runs the affiliated Fair Oil Company."),
        ("Louis & Peaches Owen Family Foundation", "Founders deceased; foundation continues East Texas health giving. Community-credibility/grant angle only."),
        ("Recent tech/company exits", "No verifiable recent Tyler exit angels found - Tyler wealth is overwhelmingly held in continuing private family enterprises."),
    ])

    _h(doc, "Flags and caveats")
    for x in [
        "CONFLICT: Jeb W. Jones is past board chair of Hospice of East Texas, the incumbent nonprofit competitor. Approach only deliberately, if at all.",
        "DILIGENCE: Dr. Kevin Pauza - an unconfirmed medical-board-discipline headline surfaced in search results; confirm before approach.",
        "All HIGH-tagged physicians need individual healthcare regulatory counsel clearance before investing. The two system CMOs (Keuer, Lamanteer) are the most sensitive - they control institutional discharge/referral pathways.",
        "Recent events in flux: Brad Brookshire retired March 2026; Curtis Mewbourne died June 2026 - family situations may be changing.",
        "Two rosters worth pulling manually (sites blocked automated access): the TEDC board page and the UT Tyler Development Board page.",
        "This document is research, not legal advice, and not an offer of securities.",
    ]:
        _p(doc, "  -  " + x)

    path = OUT + "Prospect_Research_Tyler.docx"
    doc.save(path)
    print("wrote", path)


def build_csv():
    path = OUT + "Prospect_Research_Tyler.csv"
    with open(path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["track", "rank", "name", "role_affiliation", "why_them", "aks_tag",
                    "warm_intro_path", "est_check_tier", "verification", "source_url"])
        for i, (n, role, why, pth, tier, src, ver) in enumerate(ANGELS):
            w.writerow(["angel_hnw", i + 1, n, role, why, "", pth, tier, ver, src])
        for i, (n, role, why, aks, pth, tier, src, ver) in enumerate(PHYSICIANS):
            w.writerow(["physician", i + 1, n, role, why, aks, pth, tier, ver, src])
    print("wrote", path)


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    build_docx()
    build_csv()
