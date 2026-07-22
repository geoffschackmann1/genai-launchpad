"""Draft email to Blaise Bender (sellers' counsel) transmitting the MIPA redline reply.

Companion to '17_mipa_negotiation_2026-07-22/MIPA Refuge Hospice Rev1.01 - GS Reply to
Bender Comments.docx' - the annotated redline with Geoff's threaded comment replies.
Output: sba_application/17_mipa_negotiation_2026-07-22/
"""
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docfmt
from docx import Document
from docx.shared import Pt

OUT = "sba_application/17_mipa_negotiation_2026-07-22/"


def _p(d, text, bold=False, size=11):
    p = d.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    return p


def _bullet(d, text, size=11):
    p = d.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    r.font.size = Pt(size)
    return p


def _num(d, text, size=11):
    p = d.add_paragraph(style="List Number")
    r = p.add_run(text)
    r.font.size = Pt(size)
    return p


def email_to_blaise():
    d = Document()
    d.styles["Normal"].font.name = "Aptos"
    d.styles["Normal"].font.size = Pt(11)
    d.add_heading("Draft email to Blaise Bender (sellers' counsel)", 0)
    p = d.add_paragraph()
    r = p.add_run("To: Blaise Bender | Cc: Jorge Resendiz, Dennis Hendrix | From: Geoff Schackmann | "
                  "Re: Refuge Hospice MIPA - Rev 1.01 comments, proposed final version")
    r.bold = True
    r.font.size = Pt(11)

    _p(d, "Blaise,")
    _p(d, "Thanks for turning the comments around so quickly. Attached is the updated draft with my "
          "responses to each item, organized below so you can see where things stand.")

    _p(d, "Confirmed, no open issue:", bold=True)
    for t in [
        "Buyer entity name conformed throughout (the Agreement, Exhibits, Disclosure Schedules, and "
        "signature blocks).",
        "The free and clear language restated, just cleaned up, no substantive change.",
        "Final payment terms: balloon date conformed to the Recitals, $258,489.46 stated final payment, "
        "plus the discharge, W-9, and payoff statement language.",
        "The 5 day cure period on the acceleration provision. Agreed.",
        "Exhibit C, the payment schedule, added.",
    ]:
        _bullet(d, t)

    _p(d, "Still need an answer from you:", bold=True)
    for t in [
        "Closing date. I see the “no earlier than thirty-six (36) months after the Company's CMS "
        "Certification” qualifier stays in, which protects the regulatory floor, so we're aligned on the "
        "substance. The remaining wrinkle is that “on or about January 15” leaves a window between "
        "January 8 and January 15 where the final payment figure of $258,489.46 no longer matches, since "
        "that number is computed to January 15 exactly. Simplest fix is “not before January 15, 2027.” "
        "If you want the flexibility to close in that one week window instead, the final payment needs to "
        "be restated as a formula rather than a fixed number. Either works for me, just pick one.",
        "Default rate. We're aligned on the 5 day cure, but I haven't heard back on bringing the Note's "
        "default rate down from 15% to 10%. Can we close that out too?",
        "Interim operating covenants (Section 5.03). Good to know Refuge is member managed, I'll draft "
        "consistent with that. I still need to know whether Sellers agree to the covenant in principle, "
        "and I still need the current Company Agreement so I can conform the language to it.",
        "Venue. Bexar County works for enforcement and fixes the conflict with the arbitration seat. Just "
        "confirming the arbitration seat is also Bexar, not Denton, so both provisions point to the same "
        "place.",
    ]:
        _num(d, t)

    _p(d, "Carried forward as drafted since I haven't heard back, flagging these now rather than at "
          "signing:", bold=True)
    for t in [
        "Section 1.09, Payment Deferral for Regulatory Payment Events. Gives me the ability to defer a "
        "payment if Medicare payment is delayed by a provisional period of enhanced oversight, prepayment "
        "review, or additional documentation requests. Capped at 12 months total, interest keeps accruing, "
        "no default triggered. Given Palmetto's March 2026 reactivation letter, this is a real risk, not "
        "a hypothetical one.",
        "Security agreement consolidation (Section 1.07). One joint security agreement with both Sellers, "
        "deletion of Section 2.3, and the Section 3.3-3.4 fix so the timing reflects that I don't hold the "
        "interest until Closing.",
        "Escrow for the 51% release (Section 2.03). Mirrors the security interest Sellers already hold on "
        "my 49%. I'll cover the cost.",
    ]:
        _bullet(d, t)

    _p(d, "None of these three drew a comment either way, so I want to make sure they're in front of you "
          "now instead of assumed later. If any of them is a problem, better to know now. Otherwise I'll "
          "treat them as agreed and move toward a clean version for signature.")
    _p(d, "Let me know where you land on the four items above, and flag anything on the last three.")
    _p(d, "Thanks,")
    _p(d, "Geoff", bold=True)

    docfmt.finalize(d, "Refuge Hospice MIPA - Email to Blaise Bender")
    path = OUT + "Email to Blaise Bender - MIPA Response DRAFT.docx"
    d.save(path)
    print("wrote", path)


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    email_to_blaise()
