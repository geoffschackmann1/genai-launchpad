"""Draft LOI for the acquisition of Refuge Hospice (non-binding, seller placeholders)."""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

NAVY = RGBColor(0x1F, 0x3A, 0x5F); GREY = RGBColor(0x6B, 0x6B, 0x6B)
OUT = os.path.join(os.path.dirname(__file__), "10_supporting_documents",
                   "Refuge_Hospice_LOI_DRAFT.docx")


def para(d, text, bold=False, italic=False, size=10.5, center=False, color=None):
    p = d.add_paragraph()
    if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    return p


def h(d, text):
    p = d.add_paragraph(); r = p.add_run(text)
    r.bold = True; r.font.size = Pt(11); r.font.color.rgb = NAVY
    return p


def build():
    d = Document()
    st = d.styles["Normal"].font; st.name = "Calibri"; st.size = Pt(10.5)
    for m_ in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(d.sections[0], m_, Inches(1.0))

    para(d, "LETTER OF INTENT - NON-BINDING", bold=True, size=14, center=True, color=NAVY)
    para(d, "Proposed acquisition of Refuge Hospice", center=True, size=11, color=GREY)
    para(d, "DRAFT FOR REVIEW - do not send until seller details are confirmed and counsel has reviewed.",
         italic=True, size=9, center=True, color=GREY)
    para(d, "")
    para(d, "Date: ____________, 2026")
    para(d, "To: [SELLER NAME / ENTITY], owner of Refuge Hospice ([LEGAL ENTITY NAME], a [STATE] "
            "[entity type]) (the \"Seller\" and the \"Company\")")
    para(d, "From: Tyler Hospice Hold, LLC, a Wyoming limited liability company, and/or its designated "
            "affiliate (the \"Buyer\"), by Geoff Schackmann, Manager")
    para(d, "")
    para(d, "This letter of intent (\"LOI\") sets out the principal terms on which the Buyer proposes to "
            "acquire one hundred percent (100%) of the membership interests of the Company, which holds a "
            "Medicare-certified and Medicaid-contracted hospice license serving the Tyler / East Texas "
            "market. Except for Sections 8-10 (Exclusivity, Confidentiality, and Governing Law), this LOI "
            "is non-binding and is an expression of mutual intent only.")

    h(d, "1. Transaction structure")
    para(d, "Purchase of 100% of the issued and outstanding membership interests of the Company in a "
            "change-of-ownership (CHOW) transaction, so that the Company's Medicare provider number (CCN), "
            "Medicaid contract, and state HCSSA licensure remain with the Company and convey intact. The "
            "definitive agreement will be a customary Membership Interest Purchase Agreement (\"MIPA\").")

    h(d, "2. Purchase price - $500,000 (full asking price)")
    para(d, "The aggregate purchase price is Five Hundred Thousand Dollars ($500,000), payable as follows:")
    for t in [
        "(a) $100,000 (20%) in immediately available funds at Closing (targeted July 2026);",
        "(b) $25,000 per month for four (4) consecutive months, September through December 2026 "
        "($100,000 in total); and",
        "(c) the remaining balance of $300,000, plus accrued interest under Section 3, in a single "
        "balloon payment on or before January 31, 2027 (the \"Final Payment\").",
    ]:
        para(d, t)
    para(d, "The Seller thus receives $200,000 (40% of the price) within six months of Closing and is paid "
            "in full in January 2027. The Buyer's financing for the Final Payment is already identified.")

    h(d, "3. Interest; security; guaranty")
    for t in [
        "(a) Interest. The deferred balance accrues simple interest at six percent (6.0%) per annum from "
        "Closing, paid with the Final Payment (approximately $11,000 if the schedule above is met).",
        "(b) Security. The deferred balance is secured by a pledge of the purchased membership interests "
        "of the Company until paid in full.",
        "(c) Personal guaranty. Geoff Schackmann personally and unconditionally guarantees the deferred "
        "balance.",
        "(d) Prepayment. The Buyer may prepay all or part of the balance at any time without penalty.",
    ]:
        para(d, t)

    h(d, "4. Closing; timing")
    para(d, "The parties will target a Closing within thirty (30) days of mutual acceptance of this LOI, "
            "subject to completion of due diligence and the conditions in Section 6. Time is important to "
            "both parties; the Buyer is prepared to move immediately.")

    h(d, "5. Due diligence")
    para(d, "Upon acceptance, the Seller will provide customary diligence materials, including: state "
            "licensure and survey history; Medicare and Medicaid enrollment records and any correspondence "
            "with CMS/HHSC; cost reports and cap-liability history; billing and claims history (if any); "
            "OIG/SAM exclusion confirmations for the Company and its owners; organizational documents and "
            "capitalization; tax returns; litigation, lien, and creditor information; and any employment or "
            "vendor obligations. The Buyer expects diligence to be brief given the Company's dormant status.")

    h(d, "6. Conditions to Closing")
    for t in [
        "(a) The Company's Medicare CCN, Medicaid contract, and HCSSA license are in good standing, free "
        "of revocation, suspension, deactivation, or unresolved adverse survey or enforcement action;",
        "(b) No outstanding Medicare hospice cap liability, recoupment, overpayment demand, or audit "
        "exposure (or a mutually agreed escrow/holdback covering any that exists);",
        "(c) Neither the Company nor any owner appears on the OIG LEIE or SAM.gov exclusion lists;",
        "(d) The membership interests are delivered free and clear of liens and encumbrances;",
        "(e) Regulatory notification/approval of the CHOW proceeding in the ordinary course; and",
        "(f) Execution of a mutually acceptable MIPA consistent with this LOI.",
    ]:
        para(d, t)

    h(d, "7. Conduct pending Closing")
    para(d, "From acceptance until Closing or termination, the Seller will maintain the license and "
            "enrollments in good standing, make all required filings, and not encumber, transfer, or "
            "shop the Company or its assets.")

    h(d, "8. Exclusivity (binding)")
    para(d, "For forty-five (45) days from acceptance, the Seller will negotiate exclusively with the "
            "Buyer and will not solicit, encourage, or accept any other offer for the Company or the "
            "license.")

    h(d, "9. Confidentiality (binding)")
    para(d, "Each party will keep this LOI, its terms, and information exchanged in diligence "
            "confidential, except as required by law or to professional advisors under like duty.")

    h(d, "10. Miscellaneous (binding)")
    para(d, "This LOI is governed by Texas law. Each party bears its own costs. Except for Sections 8-10, "
            "this LOI creates no binding obligation, and no obligation to consummate the transaction "
            "arises unless and until a definitive MIPA is executed by both parties.")

    para(d, "")
    para(d, "If these terms are acceptable, please countersign below, and we will circulate the MIPA and "
            "begin diligence immediately.")
    para(d, "")
    para(d, "BUYER - TYLER HOSPICE HOLD, LLC", bold=True)
    para(d, "By: ______________________________  Date: ____________")
    para(d, "Geoff Schackmann, Manager  |  geoff@azaleahospice.com")
    para(d, "")
    para(d, "ACCEPTED - SELLER", bold=True)
    para(d, "By: ______________________________  Date: ____________")
    para(d, "[SELLER NAME], [TITLE], [ENTITY]")

    d.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    build()
