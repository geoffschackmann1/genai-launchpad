"""Package audit: canonical-fact matrix + stale blacklist + formatting assertions.

Run from repo root: python3 sba_application/qa_package.py
Exit 0 only if every check passes. Writes AUDIT REPORT.docx into the package folder.
All canonical figures are the Rev 5.00 AVANT proforma outputs (QA 22/22): Avant
Hospice, LLC acquisition, $300,000 fully deferred seller note retired by the
$500,000 SBA 7(a) at month 2. Historical (already-sent, Refuge-era) drafts are
excluded from the stale scan via HISTORICAL below but still format-audited.
"""
import glob
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docfmt  # noqa: E402
import openpyxl  # noqa: E402
from docx import Document  # noqa: E402
from docx.shared import Pt  # noqa: E402
from pypdf import PdfReader  # noqa: E402

PKG = "sba_application/13_checklist_response_2026-07-16/"
PLAN = "sba_application/03_business_plan/Azalea SBA Business Plan Rev7.00 AVANT.docx"
INS = "sba_application/14_insurance_2026-07-16/"
PROFORMA = "financial_models/output/Azalea_Hospice_Proforma_Rev5.00_AVANT.xlsx"
UOP = PKG + "UOP Azalea Avant Rev3.00.xlsx"
# Refuge-era documents that were already sent (or memorialize sent correspondence).
# They legitimately reference the superseded Refuge structure, so the Avant stale
# scan skips them; formatting checks still run.
HISTORICAL = {"8 Reply Email John Mary DRAFT.docx",
              "9 Lease Summary and Flags.docx",
              "11 Reply to John Hart - Follow-up Questions DRAFT.docx"}

RESULTS = []


def check(section, name, ok, detail=""):
    RESULTS.append((section, name, bool(ok), detail))
    print(f"  {'OK ' if ok else 'FAIL'} [{section}] {name} {detail}")
    return ok


def doc_text(path):
    d = Document(path)
    t = "\n".join(p.text for p in d.paragraphs)
    for tb in d.tables:
        for row in tb.rows:
            for c in row.cells:
                t += "\n" + c.text
    for sec in d.sections:
        for p in sec.footer.paragraphs:
            t += "\n" + p.text
    return d, t


def pdf_field_text(path):
    r = PdfReader(path)
    f = r.get_fields() or {}
    return " | ".join(str(v.get("/V")) for v in f.values() if v.get("/V"))


def xlsx_text(path):
    wb = openpyxl.load_workbook(path)
    out = []
    for sh in wb.sheetnames:
        for row in wb[sh].iter_rows():
            for c in row:
                if c.value is not None:
                    out.append(str(c.value))
    return "\n".join(out)


# ---- canonical facts (Rev 5.00 AVANT) ----
REQUIRED = {
    PLAN: ["$299,195", "$695,981", "$985,429", "$130,630", "$472,185", "$733,029",
           "15.5%", "24.1%", "28.2%", "22.3%", "3.7x", "8.6x", "12.2x",
           "$6,747", "$300,000", "$303,000", "$435,000", "$250,000", "33%", "$750,000",
           "24 by month 2", "741798", "$192,358", "$1,473,592", "$225-350K", "Rev 5.00"],
    PKG + "0 Cover Email DRAFT.docx": ["$300,000", "$303,000", "$6,747", "741798"],
    PKG + "1 Checklist STATUS RESPONSE.docx": ["$300,000", "$303,000", "$250,000", "Rev 5.00", "22/22"],
    PKG + "4 Business Plan ADDENDUM Avant.docx": ["$299K", "$696K", "$985K", "$6,747", "741798",
                                                  "$225-350K", "Rev 5.00"],
    PKG + "6 Projection Assumptions Narrative.docx": ["$299K", "$696K", "$985K", "22.3%", "Rev 5.00"],
    PKG + "12 Reply to John Hart - Financing Timing and Checklist DRAFT.docx":
        ["Avant", "$300,000", "$303,000", "month 2"],
    INS + "Reply Email Todd Plummer DRAFT.docx": ["$972,167", "$276,667", "$148,000", "$170,000", "$187,500", "92-1541610", "$1,933,000"],
}
REQUIRED_XLSX = {
    UOP: ["300000", "435000", "500000", "250000", "EQUITY INJECTION", "FINANCING SEQUENCE"],
    PROFORMA: ["Working-capital policy", "ADDITIONAL FUNDING REQUIREMENT", "Revision 5.00", "Avant"],
}
# Superseded-figure blacklist: Rev 3.x figures, plus the entire Refuge / Rev 4.10
# structure ($500K price, $125K down, interim bank note $15,211/mo, $461,676 WC,
# 1/15/2027 split close, $3,354 min cash, $1,297,392 M36 cash, PTAN A91679, MIPA
# balloon $258,489.46). "A9167" and "$258,489" are plain substrings on purpose -
# under the Avant target ANY form of those references is stale in an active doc.
STALE = ["$262,951", "$680,224", "$975,925", "$100,941", "$472,928", "$746,549",
         "$78,710", "$936,793", "$924,000", "$142,524", "$182,532",
         "Rev 3.00", "Rev 3.10", "Rev3.00", "Rev3.10", "$183K", "$375K facility",
         "retire the seller balloon", "3/23/2026", "3/23/26",
         "Refuge", "Rev 4.10", "Rev4.10", "$15,211", "$461,676", "$31,250",
         "$128,491", "$457,463", "$718,153", "3.8x", "$3,354", "$1,297,392",
         "1/15/2027", "A9167", "$258,489", "interim bank note"]


def content_audit():
    print("== CONTENT: canonical facts ==")
    for path, needles in REQUIRED.items():
        _, t = doc_text(path)
        name = os.path.basename(path)
        missing = [n for n in needles if n not in t]
        check("facts", f"{name}: all {len(needles)} canonical figures present", not missing,
              f"missing {missing}" if missing else "")
    for path, needles in REQUIRED_XLSX.items():
        t = xlsx_text(path)
        name = os.path.basename(path)
        missing = [n for n in needles if n not in t]
        check("facts", f"{name}: workbook facts present", not missing,
              f"missing {missing}" if missing else "")

    print("== CONTENT: stale blacklist ==")
    targets = ([p for p in sorted(glob.glob(PKG + "*.docx"))
                if "AUDIT REPORT" not in p and os.path.basename(p) not in HISTORICAL] + [PLAN])
    for path in targets:
        _, t = doc_text(path)
        hits = [ph for ph in STALE if ph.lower() in t.lower()]
        # live-revolver scan (allow only negated phrasings)
        for m in re.finditer(r"revolver", t.lower()):
            ctx = t.lower()[max(0, m.start() - 25):m.end() + 15]
            if not any(k in ctx for k in ("no revolver", "no-revolver", "revolver or assumed", "revolver assumed")):
                hits.append("LIVE revolver")
        check("stale", os.path.basename(path), not hits, str(hits) if hits else "")
    for path in sorted(glob.glob(PKG + "*FILLED.pdf")):
        t = pdf_field_text(path)
        hits = [ph for ph in STALE if ph in t]
        check("stale", os.path.basename(path) + " (form fields)", not hits, str(hits) if hits else "")
    for path, label in [(UOP, "UOP workbook"), (PROFORMA, "proforma workbook")]:
        t = xlsx_text(path)
        needles = STALE
        if path == UOP:
            # The UOP workbook's OWN revision is 3.00 ("UOP Azalea Avant Rev3.00");
            # the "Rev 3.00" stale needles target proforma Rev 3.00 references.
            needles = [ph for ph in STALE if ph not in ("Rev 3.00", "Rev3.00")]
        hits = [ph for ph in needles if ph in t]
        check("stale", label, not hits, str(hits) if hits else "")

    print("== PDF field-fit (no clipped single-line values) ==")
    for path in sorted(glob.glob(PKG + "*FILLED.pdf")) + sorted(glob.glob(INS + "*FILLED.pdf")):
        r = PdfReader(path)
        clips = []
        for page in r.pages:
            for a in page.get("/Annots", []):
                o = a.get_object()
                if o.get("/FT") == "/Tx" and o.get("/V") and not (int(o.get("/Ff", 0)) & 4096):
                    x0, _y0, x1, _y1 = [float(v) for v in o["/Rect"]]
                    if (x1 - x0) < len(str(o["/V"])) * 4.4:
                        clips.append(str(o.get("/T")))
        check("fit", os.path.basename(path), not clips, str(clips) if clips else "")


def formatting_audit():
    print("== FORMATTING: documents ==")
    targets = [p for p in sorted(glob.glob(PKG + "*.docx"))
               if "AUDIT REPORT" not in p and "HIGHLIGHTED" not in p] + [PLAN] + sorted(glob.glob(INS + "*.docx"))
    for path in targets:
        d, t = doc_text(path)
        name = os.path.basename(path)
        xml = d.sections[0].footer._element.xml if d.sections else ""
        has_page_field = "PAGE" in xml
        heads = [p for p in d.paragraphs if p.style.name.startswith(("Heading", "Title"))]
        kwn = all(p.paragraph_format.keep_with_next for p in heads) if heads else True
        em = "—" in t
        dbl = len(re.findall(r"(?<! ) {2,}(?! )", t))
        styled = all(tb.style is not None for tb in d.tables)
        issues = []
        if not has_page_field: issues.append("no page-number field")
        if not heads: issues.append("no heading styles")
        if not kwn: issues.append("headings lack keep-with-next")
        if em: issues.append("em dash")
        if dbl > 2: issues.append(f"double spaces x{dbl}")
        if not styled: issues.append("unstyled table")
        if d.styles["Normal"].font.name != "Aptos": issues.append("Normal font not Aptos")
        for hname in ("Title", "Heading 1", "Heading 2"):
            if hname in d.styles and any(p.style.name == hname for p in d.paragraphs):
                if d.styles[hname].font.name != "Aptos Display":
                    issues.append(f"{hname} font not Aptos Display")
        check("format", name, not issues, "; ".join(issues))
    # plan-specific: TOC field + heading count
    d = Document(PLAN)
    body_xml = d.element.body.xml
    check("format", "business plan has TOC field", "TOC \\o" in body_xml or "TOC \\\\o" in body_xml)
    heads = [p for p in d.paragraphs if p.style.name.startswith("Heading")]
    check("format", "business plan heading count >= 60 (navigable)", len(heads) >= 60, f"({len(heads)})")

    print("== FORMATTING: proforma workbook ==")
    wb = openpyxl.load_workbook(PROFORMA)
    check("format", "Cover sheet is first tab", wb.sheetnames[0] == "Cover")
    frozen = sum(1 for sh in wb.sheetnames if wb[sh].freeze_panes)
    check("format", "frozen panes on tabs", frozen >= 11, f"({frozen})")
    foot = all(wb[sh].oddFooter.center.text == "Confidential" for sh in wb.sheetnames)
    check("format", "Confidential footer on every tab", foot)
    land = all(wb[sh].page_setup.orientation == "landscape" for sh in wb.sheetnames)
    check("format", "landscape print setup on every tab", land)


def write_report():
    doc = Document()
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading("Package Audit Report", 0)
    p = doc.add_paragraph()
    r = p.add_run("Azalea Hospice - SBA / lender package | Rev 5.00 AVANT basis | automated audit (sba_application/qa_package.py)")
    r.italic = True
    r.font.size = Pt(9)
    npass = sum(1 for _s, _n, ok, _d in RESULTS if ok)
    doc.add_paragraph().add_run(f"RESULT: {npass} of {len(RESULTS)} checks passed.").bold = True
    doc.add_paragraph(
        "Scope: (1) canonical-fact matrix - every financial figure cited in any active document must equal "
        "the Rev 5.00 AVANT proforma output exactly; (2) stale blacklist - figures from superseded revisions "
        "(including the entire Refuge / Rev 4.10 structure) must appear nowhere in active documents, PDF "
        "form fields, or workbooks (already-sent Refuge-era drafts are retained as history and exempted); "
        "(3) formatting - footers with live page numbers, real heading styles with keep-with-next, styled "
        "tables, no double spacing or em dashes, TOC field in the business plan, cover sheet / frozen panes "
        "/ print setup in the proforma. The proforma itself is separately verified by "
        "financial_models/qa_proforma_v3.py (22 checks: balance-sheet integrity, collections conservation, "
        "census calibration, seller-note deferral and month-2 SBA takeout, scenario toggles, hardcode scan).")
    t = doc.add_table(rows=1, cols=4)
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(["Area", "Check", "Result", "Detail"]):
        c = t.rows[0].cells[i]
        c.text = h
        for par in c.paragraphs:
            for run in par.runs:
                run.bold = True
                run.font.size = Pt(8.5)
    for sec, name, ok, detail in RESULTS:
        cells = t.add_row().cells
        for i, v in enumerate([sec, name, "PASS" if ok else "FAIL", detail]):
            cells[i].text = str(v)
            for par in cells[i].paragraphs:
                for run in par.runs:
                    run.font.size = Pt(8.5)
    docfmt.finalize(doc, "Azalea Hospice - Package Audit Report")
    path = PKG + "AUDIT REPORT.docx"
    doc.save(path)
    print("wrote", path)


if __name__ == "__main__":
    content_audit()
    formatting_audit()
    write_report()
    fails = [r for r in RESULTS if not r[2]]
    print(f"\n{len(RESULTS) - len(fails)} passed, {len(fails)} failed")
    sys.exit(1 if fails else 0)
