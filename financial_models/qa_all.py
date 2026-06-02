"""Full QA: recalc each of the three workbooks with the `formulas` engine,
scan every cell for formula errors, and spot-check the audience tabs.
"""
import os, sys, glob
import formulas
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from financial_models.engine.model import plet, NP

OUT = os.path.join(os.path.dirname(__file__), "output")
ERR_MARKERS = ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!", "#N/A", "#NULL!", "#ERROR")


def scan(path):
    fname = os.path.basename(path)
    xl = formulas.ExcelModel().loads(path).finish()
    sol = xl.calculate()
    errors = []
    cells = {}
    for key, rng in sol.items():
        if "!" not in key:
            continue
        try:
            v = rng.value
            val = v[0, 0]
        except Exception:
            continue
        cells[key] = val
        if isinstance(val, str) and any(m in val for m in ERR_MARKERS):
            errors.append((key, val))
    return fname, cells, errors


def get(cells, fname, sheet, coord):
    key = f"'[{fname}]{sheet.upper()}'!{coord}"
    return cells.get(key)


def main():
    files = {
        "SBA": "Azalea_Hospice_SBA_Loan_Package.xlsx",
        "INV": "Azalea_Hospice_Investor_Package.xlsx",
        "OPS": "Azalea_Hospice_Operations_Dashboard.xlsx",
    }
    allok = True
    for tag, fn in files.items():
        path = os.path.join(OUT, fn)
        fname, cells, errors = scan(path)
        print(f"\n===== {tag}: {fn} =====")
        print(f"  cells evaluated: {len(cells)}   formula errors: {len(errors)}")
        for e in errors[:15]:
            print("   ERR", e)
        if errors:
            allok = False
    print("\n" + ("NO FORMULA ERRORS IN ANY WORKBOOK" if allok else "ERRORS PRESENT"))

    # ---- targeted spot checks ----
    print("\n== Targeted checks ==")
    # SBA
    f = files["SBA"]; _, cells, _ = scan(os.path.join(OUT, f))
    def chk(name, val, cond):
        ok = cond(val)
        print(f"  {'OK ' if ok else 'XX '}{name}: {val}")
        return ok
    # locate S&U check + global DSCR by scanning labels is complex; read known cells
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(OUT, f))
    # find 'CHECK' row in Sources & Uses
    su = wb["Sources & Uses"]
    for row in su.iter_rows():
        for c in row:
            if isinstance(c.value, str) and "CHECK" in c.value:
                v = get(cells, f, "Sources & Uses", f"B{c.row}")
                chk("S&U Sources−Uses = 0", v, lambda x: x is not None and abs(float(x)) < 1)
    ls = wb["Lender Summary"]
    for row in ls.iter_rows():
        for c in row:
            if isinstance(c.value, str) and "Global 3-yr DSCR" in c.value:
                v = get(cells, f, "Lender Summary", f"B{c.row}")
                chk("Global 3-yr DSCR ≥ 1.25x", v, lambda x: x is not None and float(x) >= 1.25)

    # Investor
    f = files["INV"]; _, cells, _ = scan(os.path.join(OUT, f))
    wb = openpyxl.load_workbook(os.path.join(OUT, f))
    rt = wb["Returns"]
    for row in rt.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("Equity IRR"):
                v = get(cells, f, "Returns", f"B{c.row}")
                chk("Equity IRR computes (numeric)", v, lambda x: isinstance(x, (int, float)))
            if isinstance(c.value, str) and c.value.startswith("MOIC"):
                v = get(cells, f, "Returns", f"B{c.row}")
                chk("MOIC > 1.0x", v, lambda x: x is not None and float(x) > 1.0)

    # Ops
    f = files["OPS"]; _, cells, _ = scan(os.path.join(OUT, f))
    wb = openpyxl.load_workbook(os.path.join(OUT, f))
    vb = wb["Actuals vs Budget"]
    for row in vb.iter_rows():
        for c in row:
            if c.value == "Net revenue":
                v = get(cells, f, "Actuals vs Budget", f"B{c.row}")
                chk("Variance budget pulls month-3 net rev (~121k)", v, lambda x: x is not None and 100000 < float(x) < 140000)

    return allok


if __name__ == "__main__":
    ok = main()
    sys.exit(0 if ok else 1)
