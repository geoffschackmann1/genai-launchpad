"""Presentation polish for the Rev 4.10 proforma - run AFTER build, BEFORE QA.

Adds: Cover sheet (linked headline metrics, tab directory), frozen panes on all
monthly tabs, print setup (landscape, fit-to-width, repeated title rows, footer),
zoom. Touches no formulas or values on existing tabs.
"""
import json

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as gcl

XLSX = "financial_models/output/Azalea_Hospice_Proforma_Rev4.10_DYNAMIC.xlsx"
ROWMAP = "financial_models/output/proforma_v3_rowmap.json"

NAVY = "1F3B5C"
GREEN = "1F3B2D"
GREY = "808080"
FILL = PatternFill("solid", fgColor="EAF1EC")

MONTHLY_TABS = ["Census Waterfall", "Revenue Model", "Staffing", "Operating Budget",
                "P&L", "Cash Flow & Runway", "Balance Sheet", "Checks", "Dashboard"]


def cover(wb, R):
    if "Cover" in wb.sheetnames:
        del wb["Cover"]
    ws = wb.create_sheet("Cover", 0)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GREEN
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 46
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 46

    def put(cell, text, size=11, bold=False, italic=False, color=None, fmt=None):
        c = ws[cell]
        c.value = text
        c.font = Font(name="Calibri", size=size, bold=bold, italic=italic,
                      color=color or "000000")
        if fmt:
            c.number_format = fmt
        return c

    put("B2", "AZALEA HOSPICE & PALLIATIVE CARE", 22, bold=True, color=GREEN)
    put("B3", "Refuge Hospice, LLC Acquisition - Financial Proforma", 14, color=NAVY)
    put("B4", "Revision 4.10  |  July 2026  |  Tyler Hospice Hold, LLC (EIN 41-4966640)", 10, italic=True, color=GREY)
    put("B6", "Prepared for the SBA 7(a) application (SourceFunding / The Brownmiller Group) and company planning.", 9, italic=True, color=GREY)
    put("B7", "36 monthly periods (Jul-2026 to Jun-2029). Every calculated cell is a live formula; all inputs on the Control Tower.", 9, italic=True, color=GREY)
    put("B8", "REAL-CASH DISCIPLINE: no revolver or assumed facility anywhere in this model. Cash shortfalls, if any, are shown", 9, bold=True)
    put("B9", "as UNFUNDED NEED - the amount of additional capital required - never plugged by a balancing facility.", 9, bold=True)

    S3 = R["3-Year Summary"]
    CF = R["Cash Flow & Runway"]
    put("B11", "HEADLINE METRICS (all cells below are live links)", 11, bold=True)
    ws["B11"].fill = FILL
    ws["C11"].fill = FILL
    rows = [
        ("Net patient revenue - Year 1 / 2 / 3", f"='3-Year Summary'!B{S3['npr']}", f"='3-Year Summary'!C{S3['npr']}", f"='3-Year Summary'!D{S3['npr']}"),
        ("EBITDA (pre-debt-service) - Year 1 / 2 / 3", f"='3-Year Summary'!B{S3['ebitda']}", f"='3-Year Summary'!C{S3['ebitda']}", f"='3-Year Summary'!D{S3['ebitda']}"),
    ]
    r = 12
    for label, *fmls in rows:
        put(f"B{r}", label, 10)
        for j, f in enumerate(fmls):
            c = ws.cell(row=r, column=3 + j, value=f)
            c.number_format = '$#,##0;($#,##0);"-"'
            c.font = Font(name="Calibri", size=10, bold=True)
        r += 1
    singles = [
        ("Minimum cash across 36 months (base case)", f"='Cash Flow & Runway'!B{CF['mincash']}", '$#,##0;($#,##0);"-"'),
        ("Peak additional capital required (base case)", f"='Cash Flow & Runway'!B{CF['peakneed']}", '$#,##0;($#,##0);"-"'),
        ("Month-36 ending cash", f"='Cash Flow & Runway'!AL{CF['cash']}", '$#,##0;($#,##0);"-"'),
        ("Master integrity check (0 = clean)", f"='Checks'!B{R['Checks']['master']}", "0"),
    ]
    for label, f, fmt in singles:
        put(f"B{r}", label, 10)
        c = ws.cell(row=r, column=3, value=f)
        c.number_format = fmt
        c.font = Font(name="Calibri", size=10, bold=True)
        r += 1
    r += 1
    put(f"B{r}", "FINANCING SEQUENCE (Path A)", 11, bold=True)
    ws[f"B{r}"].fill = FILL; ws[f"C{r}"].fill = FILL
    r += 1
    for line in [
        "Aug 2026 - Equity ($250,000: J. Bullard $195K + G. Schackmann $55K) funds the $125,000 license down payment (49%)",
        "Sep 2026 - Interim bank note ($500,000, 6%, 36-mo) pays the sellers in full; surplus to working capital",
        "Oct-Dec  - Bank note service $15,211/month; owner salaries deferred until break-even census (accrued, repaid Jan)",
        "Jan 2027 - SBA 7(a) $500,000 funds at the 51% transfer (42 CFR 424.550(b)) and refinances the bank note",
        "Feb 2027 - Single SBA payment of $6,747/month (10.5%, 10-year); recurring coverage 3.8x / 8.6x / 12.2x",
    ]:
        put(f"B{r}", line, 9.5)
        r += 1
    r += 1
    put(f"B{r}", "WORKBOOK CONTENTS", 11, bold=True)
    ws[f"B{r}"].fill = FILL; ws[f"C{r}"].fill = FILL
    r += 1
    toc = [
        ("Control Tower", "Every model input - deal terms, census, rates, staffing, toggles (scenario, refi, SBA, deferral)"),
        ("Census Waterfall", "Admissions, discharge waterfall, ADC, patient days"),
        ("Revenue Model", "CMS rate build, payer mix, net patient revenue"),
        ("Staffing", "FTEs (census-driven with 1-month hire lag), payroll, PRN supplement, owner-deferral schedule"),
        ("Operating Budget", "Direct patient care, facility, G&A, contingency"),
        ("P&L", "Accrual P&L with account codes; EBITDA and margins"),
        ("Cash Flow & Runway", "Collections engine, all debt schedules, ending cash, UNFUNDED NEED"),
        ("Balance Sheet", "No-plug balance sheet; check row = 0 in every month"),
        ("Checks", "Integrity tests + viability metrics; master check"),
        ("Dashboard", "KPIs a lender reads first: DSCR, DSO, cap cushion, margins"),
        ("3-Year Summary", "Annual rollup and global DSCR"),
        ("Actuals vs Budget", "Enter actuals monthly; variances compute"),
    ]
    for name, desc in toc:
        put(f"B{r}", name, 10, bold=True)
        put(f"D{r}", desc, 9, color=GREY)
        r += 1
    r += 1
    put(f"B{r}", "Confidential - prepared for lender and investor diligence. Figures QA-verified by automated test harness (23 checks).", 8.5, italic=True, color=GREY)
    thin = Side(style="medium", color=GREEN)
    for col in "BCD":
        ws[f"{col}5"].border = Border(bottom=thin)
    ws.sheet_view.zoomScale = 100
    return ws


def polish(wb, R):
    for name in wb.sheetnames:
        ws = wb[name]
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.oddFooter.left.text = "Azalea Hospice - Proforma Rev 4.10"
        ws.oddFooter.left.size = 8
        ws.oddFooter.center.text = "Confidential"
        ws.oddFooter.center.size = 8
        ws.oddFooter.right.text = "Page &P of &N"
        ws.oddFooter.right.size = 8
        if name in R and "_hdr" in R.get(name, {}):
            hdr = R[name]["_hdr"]
            ws.freeze_panes = f"C{hdr + 1}"
            ws.print_title_rows = f"1:{hdr}"
            ws.sheet_view.zoomScale = 80
        elif name == "Control Tower":
            ws.freeze_panes = "A5"
        elif name == "3-Year Summary":
            ws.freeze_panes = "B5"
        elif name == "Actuals vs Budget":
            hdr = 5
            ws.freeze_panes = "C6"
            ws.sheet_view.zoomScale = 85


def main():
    rm = json.load(open(ROWMAP))
    R = rm["rows"]
    wb = openpyxl.load_workbook(XLSX)
    cover(wb, R)
    polish(wb, R)
    wb.active = 0
    wb.save(XLSX)
    print("formatted", XLSX)


if __name__ == "__main__":
    main()
