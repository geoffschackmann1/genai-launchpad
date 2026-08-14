"""Build a workbook, recalc it with the pure-Python `formulas` engine, then
compare the recalculated cell values against the Python reference model
(model.compute). LibreOffice is unavailable in this environment.
"""
import os, sys
import formulas
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from financial_models.engine.build import Book, build_engine
from financial_models.engine.model import compute, PERIODS, NP, pcol, plet
from openpyxl.utils import get_column_letter

OUT = os.path.join(os.path.dirname(__file__), "output")
os.makedirs(OUT, exist_ok=True)


def build_test():
    bk = Book("ENGINE TEST")
    build_engine(bk)
    bk.wb.calculation.fullCalcOnLoad = True
    p = os.path.join(OUT, "_engine_test.xlsx")
    bk.wb.save(p)
    return bk, p


class Recalc:
    """Wrap the `formulas` solution for cell lookups by (sheet,row,col)."""
    def __init__(self, path):
        self.fname = os.path.basename(path)
        xl = formulas.ExcelModel().loads(path).finish()
        self.sol = xl.calculate()

    def cell(self, sheet, row, col_letter):
        key = f"'[{self.fname}]{sheet.upper()}'!{col_letter}{row}"
        if key not in self.sol:
            return None
        v = self.sol[key].value
        try:
            return float(v[0, 0])
        except Exception:
            try:
                return v[0, 0]
            except Exception:
                return v


def main():
    bk, p = build_test()
    print("Built", p)
    rc = Recalc(p)
    R, a = compute()

    checks = []
    def cmp(name, sheet, rowkey, arr, tol=1.0, idxs=None):
        rows = bk.rows[sheet]
        idxs = idxs or range(NP)
        worst = 0; worst_i = -1
        for i in idxs:
            v = rc.cell(sheet, rows[rowkey], plet(i))
            v = 0.0 if v is None else float(v)
            d = abs(v - arr[i])
            if d > worst: worst = d; worst_i = i
        ok = worst <= tol
        checks.append(ok)
        flag = "OK " if ok else "XX "
        print(f"  {flag}{name:<34} max|Δ|={worst:12.4f}  (worst {PERIODS[worst_i]['label']})")

    print("\n== Revenue ==")
    cmp("Patient-days", "Revenue Model", "pd", R["pd"], 0.1)
    cmp("Gross revenue", "Revenue Model", "gross", R["gross"])
    cmp("Net revenue", "Revenue Model", "net", R["net"])
    print("== Staffing ==")
    cmp("Direct-care labor (COGS)", "Staffing & Payroll", "cogs_labor", R_cogs_labor(R))
    cmp("Indirect labor (SG&A)", "Staffing & Payroll", "sga_labor", R_sga_labor(R))
    print("== P&L ==")
    cmp("EBITDA", "Operating Budget", "ebitda", R["ebitda"])
    cmp("Net income", "Operating Budget", "ni", R["ni"])
    print("== Debt ==")
    cmp("SBA interest", "Debt Schedule", "int", R["int_sba"])
    cmp("SBA ending balance", "Debt Schedule", "end", R["sba_bal"])
    print("== Cash Flow & BS ==")
    cmp("Ending cash", "Cash Flow & BS", "endcash", R["end_cash"], tol=2.0)
    cmp("LOC balance", "Cash Flow & BS", "locbal", R["loc_bal"], tol=2.0)

    # Balance-sheet check row must be ~0 every period (incl. opening col B)
    rows = bk.rows["Cash Flow & BS"]; bschk = rows["bs_check"]
    worst = 0
    cols = ["B"] + [plet(i) for i in range(NP)]
    for cl in cols:
        v = rc.cell("Cash Flow & BS", bschk, cl)
        worst = max(worst, abs(float(v or 0)))
    ok = worst < 1.0
    checks.append(ok)
    print(f"\n  {'OK ' if ok else 'XX '}Balance sheet CHECK row  max|Assets-(L+E)|={worst:.6f}")

    print("\n" + ("ALL CHECKS PASSED" if all(checks) else "FAILURES PRESENT"))
    return all(checks)


def R_cogs_labor(R):
    return [R["ft_direct"][i] + R["prn"][i] + R["med_dir"][i] + R["burden_d"][i] + R["health_d"][i] for i in range(NP)]

def R_sga_labor(R):
    return [R["ft_indirect"][i] + R["rhonda"][i] + R["burden_i"][i] + R["health_i"][i] for i in range(NP)]


if __name__ == "__main__":
    ok = main()
    sys.exit(0 if ok else 1)
